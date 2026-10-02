# -*- coding: utf-8 -*-
"""
Converte as 12 bases mensais .xlsx do lab em .csv com cara de exportacao bruta
(estilo "Salvar como CSV" de um sistema pt-BR), preservando os valores atuais.

Uso:
    python converter_xlsx_para_csv_bruto.py                 # gera e verifica os CSV
    python converter_xlsx_para_csv_bruto.py --remover-xlsx  # idem e, se tudo passar, remove os .xlsx da pasta

Regras (ver plano do lab):
  - separador ';', UTF-8 com BOM, CRLF, quoting minimo
  - texto: escrito exatamente como esta na celula
  - numero tipado: 15 digitos significativos (como o Excel exibe), decimal com virgula
  - data tipada: dd/mm/yyyy ; datetime: dd/mm/yyyy HH:MM:SS (arredondado ao segundo)
  - preambulo de metadados (4 linhas fixas + opcionais por mes) + linha em branco + cabecalho
  - rodape: linha em branco + "Total de registros;N" + "Fim da exportacao"
"""
import csv
import datetime as dt
import glob
import os
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

AQUI = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.dirname(AQUI)
BASES = os.path.join(LAB, "01_bases_mensais")
BACKUP = os.path.join(LAB, "_backup_xlsx_originais")

TITULO = "Relatorio de assinaturas - exportacao bruta"
LINHAS_OPCIONAIS = {
    "usuario": ["Usuario", "integracao.bi"],
    "filtro": ["Filtro", "status <> 'excluida'"],
    "fuso": ["Fuso horario", "America/Sao_Paulo"],
}
# Quais linhas opcionais entram em cada mes (deterministico). 4 fixas + estas.
OPCIONAIS_POR_MES = {
    "01": [],
    "02": ["usuario"],
    "03": [],
    "04": ["usuario", "filtro"],
    "05": ["filtro"],
    "06": ["usuario", "filtro", "fuso"],
    "07": [],
    "08": ["fuso"],
    "09": ["usuario", "fuso"],
    "10": [],
    "11": ["usuario", "filtro", "fuso"],
    "12": ["usuario"],
}
RODAPE_FIM = "Fim da exportacao"


def fmt(v):
    """Serializa um valor de celula como um export pt-BR faria."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, dt.datetime):
        v = (v + dt.timedelta(milliseconds=500)).replace(microsecond=0)
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        s = format(v, ".15g")
        if "e" in s or "E" in s:
            s = format(v, "f").rstrip("0").rstrip(".")
        return s.replace(".", ",")
    return str(v)


def ler_metadata(wb):
    ws = wb["metadata_exportacao"]
    meta = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] and row[0] != "campo":
            meta[str(row[0])] = row[1] if len(row) > 1 else None
    return meta


def ultima_linha_xml(caminho):
    """Numero da ultima <row> da aba 'dados' lido direto do XML (independente do openpyxl)."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with zipfile.ZipFile(caminho) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        for s in wb.find("m:sheets", ns):
            if s.get("name") != "dados":
                continue
            alvo = rels[s.get("{%s}id" % ns["r"])].lstrip("/")
            alvo = alvo if alvo.startswith("xl/") else "xl/" + alvo
            ws = ET.fromstring(z.read(alvo))
            ultimo = 0
            for r in ws.find("m:sheetData", ns):
                ultimo = max(ultimo, int(r.get("r")))
            return ultimo
    raise RuntimeError("aba 'dados' nao encontrada em " + caminho)


def converter(caminho_xlsx):
    nome = os.path.basename(caminho_xlsx)
    mes = re.search(r"_2025_(\d{2})_", nome).group(1)
    caminho_csv = os.path.splitext(caminho_xlsx)[0] + ".csv"

    wb = openpyxl.load_workbook(caminho_xlsx, read_only=True, data_only=True)
    meta = ler_metadata(wb)
    ws = wb["dados"]

    preambulo = [
        [TITULO],
        ["Competencia", fmt(meta.get("competencia"))],
        ["Sistema origem", fmt(meta.get("sistema_origem"))],
        ["Gerado em", fmt(meta.get("gerado_em"))],
    ]
    for chave in OPCIONAIS_POR_MES[mes]:
        preambulo.append(list(LINHAS_OPCIONAIS[chave]))
    linha_cabecalho = len(preambulo) + 2  # +1 linha em branco, +1 = a propria linha do cabecalho

    cabecalho = None
    amostra_original = []  # (indice_dados, tupla_original) para verificacao
    n_dados = 0
    with open(caminho_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        for linha in preambulo:
            w.writerow(linha)
        w.writerow([])
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            row = list(row)
            if i == 0:
                # cabecalho: corta apenas celulas vazias a direita
                while row and (row[-1] is None or row[-1] == ""):
                    row.pop()
                cabecalho = [str(c) for c in row]
                w.writerow(cabecalho)
                continue
            ncol = len(cabecalho)
            if len(row) > ncol:
                extra = row[ncol:]
                if any(v is not None and v != "" for v in extra):
                    raise RuntimeError(f"{nome}: linha {i + 1} tem valor fora do cabecalho: {extra}")
                row = row[:ncol]
            elif len(row) < ncol:
                row = row + [None] * (ncol - len(row))
            w.writerow([fmt(v) for v in row])
            n_dados += 1
            amostra_original.append((n_dados, tuple(row)))
        w.writerow([])
        w.writerow(["Total de registros", str(n_dados)])
        w.writerow([RODAPE_FIM])
    wb.close()

    return {
        "nome": nome, "mes": mes, "csv": caminho_csv, "xlsx": caminho_xlsx,
        "linha_cabecalho": linha_cabecalho, "cabecalho": cabecalho,
        "n_dados": n_dados, "amostra": amostra_original,
        "ultima_linha_xml": ultima_linha_xml(caminho_xlsx),
    }


def verificar(info):
    erros = []
    with open(info["csv"], "rb") as f:
        if f.read(3) != b"\xef\xbb\xbf":
            erros.append("sem BOM UTF-8")
    with open(info["csv"], "r", encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    lc = info["linha_cabecalho"]
    if linhas[lc - 1] != info["cabecalho"]:
        erros.append(f"cabecalho na linha {lc} difere do original: {linhas[lc - 1][:5]}...")
    if linhas[lc - 2] != []:
        erros.append("linha antes do cabecalho nao esta em branco")
    dados = linhas[lc: lc + info["n_dados"]]
    ncol = len(info["cabecalho"])
    ruins = [k for k, r in enumerate(dados) if len(r) != ncol]
    if ruins:
        erros.append(f"{len(ruins)} linhas de dados com numero de campos != {ncol} (ex.: {ruins[:3]})")
    rodape = linhas[lc + info["n_dados"]:]
    if rodape != [[], ["Total de registros", str(info["n_dados"])], [RODAPE_FIM]]:
        erros.append(f"rodape inesperado: {rodape}")
    esperado = info["ultima_linha_xml"] - 1
    if info["n_dados"] != esperado:
        erros.append(f"linhas de dados {info['n_dados']} != ultima linha do XML - 1 = {esperado}")
    # amostra: primeira, ultima e 3 intermediarias, campo a campo
    idx = sorted(set([1, info["n_dados"]] + [max(1, info["n_dados"] * k // 4) for k in (1, 2, 3)]))
    amostra = dict(info["amostra"])
    for i in idx:
        orig = amostra[i]
        lido = dados[i - 1]
        for c, (vo, vl) in enumerate(zip(orig, lido)):
            if isinstance(vo, str):
                if vo != vl:
                    erros.append(f"linha dados {i} col {c} texto difere: {vo!r} vs {vl!r}")
            elif fmt(vo) != vl:
                erros.append(f"linha dados {i} col {c} valor difere: {vo!r} -> {fmt(vo)!r} vs {vl!r}")
    return erros


def main():
    remover = "--remover-xlsx" in sys.argv
    arquivos = sorted(glob.glob(os.path.join(BASES, "saas_assinaturas_2025_*.xlsx")))
    if len(arquivos) != 12:
        print(f"AVISO: esperados 12 xlsx, encontrados {len(arquivos)}")
    resultados = []
    tudo_ok = True
    print(f"{'arquivo':<40} {'cab.':>4} {'dados':>6} {'cols':>4} {'tam KB':>7}  status")
    for p in arquivos:
        info = converter(p)
        erros = verificar(info)
        ok = not erros
        tudo_ok &= ok
        tam = os.path.getsize(info["csv"]) // 1024
        print(f"{os.path.basename(info['csv']):<40} {info['linha_cabecalho']:>4} {info['n_dados']:>6} "
              f"{len(info['cabecalho']):>4} {tam:>7}  {'OK' if ok else 'ERRO'}")
        for e in erros:
            print("     - " + e)
        resultados.append(info)

    if not tudo_ok:
        print("\nHouve erros de verificacao. Nenhum .xlsx foi removido.")
        sys.exit(1)
    print("\nTodos os 12 CSV gerados e verificados.")

    if remover:
        for info in resultados:
            bkp = os.path.join(BACKUP, info["nome"])
            if not (os.path.exists(bkp) and os.path.getsize(bkp) == os.path.getsize(info["xlsx"])):
                print(f"Backup ausente/diferente para {info['nome']}; .xlsx NAO removido.")
                sys.exit(2)
        for info in resultados:
            os.remove(info["xlsx"])
        print(f"{len(resultados)} arquivos .xlsx removidos de 01_bases_mensais (backup em _backup_xlsx_originais).")


if __name__ == "__main__":
    main()
